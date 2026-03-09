import os
import openpyxl

DATA_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "data", "atm_data.xlsx")

def clear_disputes():
    try:
        if not os.path.exists(DATA_FILE):
            print(f"[ERROR] Data file not found at {DATA_FILE}")
            return

        wb = openpyxl.load_workbook(DATA_FILE)
        if "disputes" in wb.sheetnames:
            ws = wb["disputes"]
            max_r = ws.max_row
            
            # Keep row 1 (headers), delete rest
            if max_r > 1:
                ws.delete_rows(2, max_r)
                wb.save(DATA_FILE)
                print(f"[SUCCESS] Removed {max_r - 1} entries from the 'disputes' sheet.")
            else:
                print("[INFO] No entries to clear. The 'disputes' sheet is already empty.")
        else:
            print("[ERROR] 'disputes' sheet not found in the Excel file.")
    except PermissionError:
        print(f"[ERROR] Permission Denied. Please ensure 'atm_data.xlsx' is CLOSED in Excel before running this.")
    except Exception as e:
        print(f"[ERROR] Failed to clear disputes: {e}")

if __name__ == "__main__":
    clear_disputes()
