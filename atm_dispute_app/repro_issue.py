
import os
import openpyxl
from openpyxl import Workbook

# Mock constants/logic from ATM_DISPUTE.py
DISPUTES_SHEET_NAME = "disputes"
headers = ["ref", "txndate", "cardno", "acno", "atmid", "txnno", "amount", "branch",
           "debit_ac", "credit_ac", "status", "posting_date", "file_type",
           "remarks", "posting_flag", "posting_user", "card_fiid", "card_bank_name",
           "term_fiid", "term_bank_name", "comp_type", "disp_type", "src_ip", "user_name"]

# Mock Record Data
ref = "REF001"
txndate_fmt = "01-JAN-25"
cardno = "1111222233334444"
acno = "0000012345678"
atmid = "S10A0001"
txnno = "TX100"
amount = "500"
branch = "00100"
debit_ac = "2399724042928" # BGL
credit_ac = "0000012345678" # Cust
file_type = "CR"
idx = 1
user_name = "TESTUSER"
card_fiid = "C001"
card_bank_name = "SBIG"
term_fiid = "F001"
term_bank_name = "HDFC"
comp_type = "SOF"
disp_type = "short"

record = [
    ref, txndate_fmt, cardno, acno, atmid, txnno, amount, branch,
    debit_ac, credit_ac, "NO", None, file_type,  f"File{idx} Entry",
    "NO", user_name, card_fiid, card_bank_name,
    term_fiid, term_bank_name, comp_type, disp_type, "127.0.0.1", user_name 
]

filename = "test_data.xlsx"

# 1. Create File
wb = Workbook()
if 'Sheet' in wb.sheetnames: wb.remove(wb['Sheet'])
wb.create_sheet(DISPUTES_SHEET_NAME)
ws = wb[DISPUTES_SHEET_NAME]
ws.append(headers)
ws.append(record)
wb.save(filename)
print("Created test_data.xlsx")

# 2. Read back using main.py logic
print("\n--- Reading back with main.py logic ---")
wb2 = openpyxl.load_workbook(filename)
sheet = wb2[DISPUTES_SHEET_NAME]

for row in sheet.iter_rows(min_row=2, values_only=True):
    print(f"Row Length: {len(row)}")
    
    main_debit = row[8]
    main_credit = row[9]
    main_status = row[10]
    
    print(f"Index 8 (Debit):  Expected='{debit_ac}', Got='{main_debit}'")
    print(f"Index 9 (Credit): Expected='{credit_ac}', Got='{main_credit}'")
    print(f"Index 10 (Status): Expected='NO', Got='{main_status}'")
    
    if main_debit != debit_ac:
        print("!! DEBIT MISMATCH !!")
    if main_credit != credit_ac:
        print("!! CREDIT MISMATCH !!")
    if main_status != "NO":
        print("!! STATUS MISMATCH !!")
        
    # Check shift
    if main_debit == "NO":
        print("ALERT: Debit column contains 'NO'. Data is shifted LEFT by 2 columns?")
    if main_credit == "NO":
        print("ALERT: Credit column contains 'NO'. Data is shifted LEFT by 1 column?")

wb2.close()
