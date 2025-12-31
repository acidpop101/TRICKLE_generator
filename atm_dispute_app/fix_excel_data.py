import openpyxl
import os

DATA_FILE = r"d:\TRICKLE_generator\atm_dispute_app\atm_data.xlsx"

# Copy of Constants from constants.py
checkDigitConstantArray = [
    [10, 9, 8, 7, 6, 5, 4, 3, 2],
    [5, 10, 4, 9, 3, 8, 2, 7, 1],
    [8, 5, 2, 10, 7, 4, 1, 9, 6],
    [4, 8, 1, 5, 9, 2, 6, 10, 3],
    [2, 4, 6, 8, 10, 1, 3, 5, 7],
    [1, 2, 3, 4, 5, 6, 7, 8, 9],
    [6, 1, 7, 2, 8, 3, 9, 4, 10],
    [3, 6, 9, 1, 4, 7, 10, 2, 5],
    [7, 3, 10, 6, 2, 9, 5, 1, 8],
    [9, 7, 5, 3, 1, 10, 8, 6, 4],
    [10, 9, 8, 7, 6, 5, 4, 3, 2],
    [5, 10, 4, 9, 3, 8, 2, 7, 1],
    [8, 5, 2, 10, 7, 4, 1, 9, 6],
    [4, 8, 1, 5, 9, 2, 6, 10, 3],
    [2, 4, 6, 8, 10, 1, 3, 5, 7],
    [1, 2, 3, 4, 5, 6, 7, 8, 9]
]

def get_check_digit(base_acno_str):
    """Calculate check digit for account number string."""
    try:
        Macno = int(base_acno_str)
    except:
        return "0" # Fallback
        
    JE = 15
    Mchkdigit = 0
    Cdigit = 0

    while Macno > 0:
        iLastDigit = Macno % 10
        if iLastDigit > 0:
            Mdigit = iLastDigit - 1
            indexI = JE
            indexJ = Mdigit
            
            if 0 <= indexI < 16 and 0 <= indexJ < 9:
                Mchkdigit += checkDigitConstantArray[indexI][indexJ]
    
        Macno //= 10
        JE -= 1

    Cdigit = Mchkdigit % 10
    return str(Cdigit)

def fix_excel_data():
    if not os.path.exists(DATA_FILE):
        print(f"File {DATA_FILE} not found.")
        return

    wb = openpyxl.load_workbook(DATA_FILE)
    
    # 1. Load Account Details
    AC_LIST_SHEET = "ac_details"
    if AC_LIST_SHEET not in wb.sheetnames:
        print(f"Sheet {AC_LIST_SHEET} not found.")
        return
        
    ac_sheet = wb[AC_LIST_SHEET]
    fiid_map = {}
    col_fiid, col_vostro, col_settl = 0, 1, 2
    for row in ac_sheet.iter_rows(min_row=2, values_only=True):
        if not row[col_fiid]: continue
        fiid = str(row[col_fiid]).strip()
        vostro = str(row[col_vostro]).strip() if row[col_vostro] else ""
        settl = str(row[col_settl]).strip() if row[col_settl] else ""
        fiid_map[fiid] = {"settl": settl, "vostro": vostro}

    # 2. Fix Disputes Sheet
    target_sheet_name = "disputes"
    sheet = wb[target_sheet_name]

    # Indices (1-BASED) using verified hardcoded values
    IDX_COMP = 21
    IDX_DISP = 22
    IDX_CARD_FIID = 17
    IDX_FILE_TYPE = 13
    IDX_BRANCH = 8
    IDX_DEBIT = 9
    IDX_CREDIT = 10

    count = 0
    for row in sheet.iter_rows(min_row=2):
        try:
            comp_type = str(row[IDX_COMP-1].value).strip() if row[IDX_COMP-1].value else ""
            disp_type = str(row[IDX_DISP-1].value).strip() if row[IDX_DISP-1].value else ""
            
            if comp_type == "FOS" and disp_type == "short":
                card_fiid = row[IDX_CARD_FIID-1].value
                file_type = row[IDX_FILE_TYPE-1].value
                br_val = row[IDX_BRANCH-1].value
                branch = str(br_val).strip().zfill(5) if br_val else "00000"
                
                ac_info = fiid_map.get(card_fiid)
                if not ac_info: continue
                    
                settl_ac = ac_info["settl"]
                vostro_ac = ac_info["vostro"]
                
                if file_type == "T1":
                    # Logic: 98581 + Branch + CheckDigit
                    base_ac = f"98581{branch}"
                    cd = get_check_digit(base_ac)
                    new_dr = f"{base_ac}{cd}"
                    new_cr = settl_ac
                    
                    row[IDX_DEBIT-1].value = new_dr
                    row[IDX_CREDIT-1].value = new_cr
                    print(f"Fixed Row (T1): Dr={new_dr}, Cr={new_cr}")
                    count += 1
                    
                elif file_type == "T2":
                    new_dr = settl_ac
                    new_cr = vostro_ac
                    row[IDX_DEBIT-1].value = new_dr
                    row[IDX_CREDIT-1].value = new_cr
                    print(f"Fixed Row (T2): Dr={new_dr}, Cr={new_cr}")
                    count += 1
        except IndexError:
            continue

    wb.save(DATA_FILE)
    print(f"SUCCESS: Fixed {count} rows in {DATA_FILE}")

if __name__ == "__main__":
    fix_excel_data()
