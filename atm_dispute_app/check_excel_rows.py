
import openpyxl
from openpyxl import load_workbook

FILE_PATH = r"d:\TRICKLE_generator\atm_dispute_app\atm_data.xlsx"

def check_rows():
    wb = load_workbook(FILE_PATH)
    sheet = wb["disputes"]
    
    print(f"{'Row':<5} | {'Txn':<5} | {'File':<5} | {'Debit Ac':<15} | {'Credit Ac':<15}")
    print("-" * 60)
    
    for i, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        txn = str(row[5]).strip() if row[5] else ""
        if txn == "2140":
            file_type = str(row[12])
            dr = str(row[8])
            cr = str(row[9])
            print(f"{i:<5} | {txn:<5} | {file_type:<5} | {dr:<15} | {cr:<15}")

if __name__ == "__main__":
    check_rows()
