import os
import openpyxl
from datetime import datetime

PATH = r"d:\TRICKLE_generator\atm_dispute_app\atm_data.xlsx"
SHEET = "disputes"

def restore():
    try:
        if os.path.exists(PATH):
            wb = openpyxl.load_workbook(PATH)
        else:
            wb = openpyxl.Workbook()
            # Remove default sheet if needed, or rename
            if "Sheet" in wb.sheetnames:
                del wb["Sheet"]
                
        if SHEET not in wb.sheetnames:
            print("Sheet not found, creating...")
            ws = wb.create_sheet(SHEET)
        else:
            ws = wb[SHEET]
        
        # Clear existing data just in case
        ws.delete_rows(2, ws.max_row)
        
        # Define 6 rows
        # Cols: Ref(0), TxnDate(1), Card(2), Ac(3), ATM(4), Txn(5), Amt(6), Br(7), Dr(8), Cr(9), Stat(10), PDate(11), Type(12)
        # 1-based indexing for openpyxl
        
        rows = [
            # Pair 1: T1
            ["1", "06-Oct-25", "4359786001519159", "0", "T1BW000177031", "0880", "21000", "00177", "98581001774", "4897928042921", "O", "06-Oct-25", "T1"],
            # Pair 1: T2 (Dr=Settlement 48979..., Cr=Vostro 30118...)
            ["1", "06-Oct-25", "4359786001519159", "0", "T1BW000177031", "0880", "21000", "00177", "4897928042921", "30118760239", "O", "06-Oct-25", "T2"],
            
            # Pair 2: T1
            ["2", "21-Nov-25", "4359786000423874", "0", "T1BS020134107", "3667", "21000", "20134", "98581201344", "4897928042921", "O", "21-Nov-25", "T1"],
            # Pair 2: T2
            ["2", "21-Nov-25", "4359786000423874", "0", "T1BS020134107", "3667", "21000", "20134", "4897928042921", "30118760239", "O", "21-Nov-25", "T2"],
            
            # Pair 3: T1
            ["3", "21-Dec-25", "4359786000300866", "0", "T1BW000138136", "2140", "21000", "00138", "98581001388", "4897928042921", "O", "21-Dec-25", "T1"],
            # Pair 3: T2
            ["3", "21-Dec-25", "4359786000300866", "0", "T1BW000138136", "2140", "21000", "00138", "4897928042921", "30118760239", "O", "21-Dec-25", "T2"]
        ]
        
        for r_idx, row_data in enumerate(rows, 2):
            for c_idx, val in enumerate(row_data, 1):
                ws.cell(row=r_idx, column=c_idx, value=val)
                
        wb.save(PATH)
        print("Restored 6 rows.")
        
    except Exception as e:
        print(f"Error: {e}")

if __name__ == "__main__":
    restore()
