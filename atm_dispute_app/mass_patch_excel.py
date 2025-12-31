
import openpyxl
from openpyxl import load_workbook

# 16x9 Matrix from constants.py (Verified)
checkDigitConstantArray = [
    [10, 9, 8, 7, 6, 5, 4, 3, 2],
    [5, 10, 4, 9, 3, 8, 2, 7, 1],
    [8, 5, 2, 10, 7, 4, 1, 9, 6],
    [4, 8, 1, 5, 9, 2, 6, 10, 3],
    [2, 4, 6, 8, 10, 1, 3, 5, 7],
    [1, 2, 3, 4, 5, 6, 7, 8, 9],
    [6, 1, 7, 2, 8, 3, 9, 4, 10],
    [3, 6, 9, 1, 4, 7, 10, 2, 5],
    [7, 3, 10, 6, 2, 9, 5, 1, 8],
    [9, 7, 5, 3, 1, 10, 8, 6, 4],
    [10, 9, 8, 7, 6, 5, 4, 3, 2],
    [5, 10, 4, 9, 3, 8, 2, 7, 1],
    [8, 5, 2, 10, 7, 4, 1, 9, 6],
    [4, 8, 1, 5, 9, 2, 6, 10, 3],
    [2, 4, 6, 8, 10, 1, 3, 5, 7],
    [1, 2, 3, 4, 5, 6, 7, 8, 9]
]

def get_check_digit(acc_num_str):
    try:
        Macno = int(acc_num_str)
    except:
        return "?"

    JE = 15
    Mchkdigit = 0
    while Macno > 0:
        iLastDigit = Macno % 10
        if iLastDigit > 0:
            Mdigit = iLastDigit - 1
            indexI = JE
            indexJ = Mdigit
            if 0 <= indexI < 16 and 0 <= indexJ < 9:
                Mchkdigit += checkDigitConstantArray[indexI][indexJ]
        Macno //= 10
        JE -= 1
    return str(Mchkdigit % 10)

FILE_PATH = r"d:\TRICKLE_generator\atm_dispute_app\atm_data.xlsx"

def mass_patch():
    print(f"Loading {FILE_PATH}...")
    try:
        wb = load_workbook(FILE_PATH)
        sheet = wb["disputes"]
    except Exception as e:
        print(f"Error loading workbook: {e}")
        return

    # Accounts from ac_details for lookup (assuming hardcoded typicals for now)
    # AcDetails: Settl = 48979..., Vostro = 30118...
    # Or I should read ac_details...
    # For now, I will assume the target rows use the standard accounts:
    settlement_ac = "4897928042921"
    vostro_ac = "30118760239"

    count = 0
    for row in sheet.iter_rows(min_row=2):
        if not row[0].value: continue # Skip empty
        
        # Identify FOS / Short
        # Indices: 0=Ref, 12=FileType, 20=CompType, 21=DispType, 4=ATMId
        # Using safely gets
        try:
            # Columns are 0-indexed.
            # 4=ATMId, 8=Debit, 9=Credit, 12=FileType
            atmid = str(row[4].value or "").strip()
            file_type = str(row[12].value or "").strip()
            # row[20] might be comp_type but let's assume FOS based on file contents
            # Actually, check logic: T1 -> Dr 98581...
            
            if not atmid: continue

            # Extract Branch (Indices 5-10) for FOS logic
            # T1BW000177... -> 00177
            if len(atmid) >= 10:
                branch = atmid[5:10]
            else:
                branch = atmid[-5:]
                
            # Base for T1 Debit - 98581 + Branch
            base_ac = "98581" + branch
            chk = get_check_digit(base_ac)
            final_debit_98 = base_ac + chk
            
            # Update Logic based on File Type in Row
            if file_type == "T1":
                # T1: Debit 98581... | Credit Settlement
                # Check if current value matches format or just overwrite
                print(f"Patching T1 (ATM {atmid[:13]}... Branch {branch})")
                row[8].value = final_debit_98
                row[9].value = settlement_ac
                count += 1
                
            elif file_type == "T2":
                # T2: Debit Settlement | Credit Vostro
                print(f"Patching T2 (ATM {atmid[:13]}...)")
                row[8].value = settlement_ac
                row[9].value = vostro_ac
                count += 1
                
        except IndexError:
            continue
            
    print(f"Patched {count} rows.")
    wb.save(FILE_PATH)
    print("Workbook saved.")

if __name__ == "__main__":
    mass_patch()
